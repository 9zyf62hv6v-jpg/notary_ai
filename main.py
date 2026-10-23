import json
import re
import os
import time
import smtplib
import urllib3
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from fastapi import FastAPI, Request, Form, File, UploadFile
from fastapi.templating import Jinja2Templates
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse, FileResponse
from fractions import Fraction
import pytesseract
from PIL import Image
import io
from pillow_heif import register_heif_opener
import sqlite3
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import requests
from bs4 import BeautifulSoup
from duckduckgo_search import DDGS

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

register_heif_opener()
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

app = FastAPI(title="Нотариальный ИИ-Ассистент")
templates = Jinja2Templates(directory="templates")

if not os.path.exists("downloads"):
    os.makedirs("downloads")

last_checker_result = {"text": ""}

def init_db():
    conn = sqlite3.connect('notary.db')
    cursor = conn.cursor()
    cursor.execute('''CREATE TABLE IF NOT EXISTS clients (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            full_name TEXT, series TEXT, number TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''')
    try:
        cursor.execute("ALTER TABLE clients ADD COLUMN phone TEXT")
        cursor.execute("ALTER TABLE clients ADD COLUMN email TEXT")
    except sqlite3.OperationalError:
        pass
    cursor.execute('''CREATE TABLE IF NOT EXISTS cases (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            client_id INTEGER, case_number TEXT, description TEXT,
            status TEXT DEFAULT 'В работе', created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(client_id) REFERENCES clients(id))''')
    conn.commit()
    conn.close()

init_db()

def clean_ai_text(text):
    if not text: return ""
    text = re.sub(r'[\u4e00-\u9fff]+', '', text)
    text = text.replace('```', '').replace('**', '').replace('*', '').replace('---', '')
    lines = text.split('\n')
    clean_lines = []
    for line in lines:
        if len(re.findall(r'[a-zA-Z]', line)) > len(re.findall(r'[а-яА-Я]', line)) and len(re.findall(r'[a-zA-Z]', line)) > 5:
            continue
        clean_lines.append(line)
    text = '\n'.join(clean_lines)
    text = re.sub(r'\[.*?\]', '___________________', text)
    result = text.replace('"]', '').replace('".', '').replace("myself", "себя").replace(" my ", " мой ")
    return result.strip()

# --- УМНЫЙ ПОИСК (С ЗАПАСНЫМ ВАРИАНТОМ) ---
def search_web(query):
    # 1. Проверка на курс валют
    if "курс" in query.lower() and ("доллар" in query.lower() or "usd" in query.lower() or "евро" in query.lower()):
        try:
            resp = requests.get("https://www.cbr-xml-daily.ru/daily_json.js", timeout=5)
            data = resp.json()
            return f"📊 Официальный курс ЦБ РФ на сегодня:\nДоллар: {data['Valute']['USD']['Value']} руб.\nЕвро: {data['Valute']['EUR']['Value']} руб."
        except:
            pass

    # 2. Чистим запрос от слов-паразитов
    clean_query = query.lower()
    for word in ["выдай", "последнее обновление", "найди", "узнай", "на сайте", "прочитай", "пожалуйста", "какой", "сегодня", "агент"]:
        clean_query = clean_query.replace(word, "")
    clean_query = clean_query.strip()
    if not clean_query: clean_query = query

    # 3. Пробуем DuckDuckGo (HTML версия)
    try:
        url = f"https://html.duckduckgo.com/html/?q={requests.utils.quote(clean_query)}"
        headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
        resp = requests.get(url, headers=headers, timeout=10, verify=False)
        soup = BeautifulSoup(resp.text, 'html.parser')
        
        results = soup.find_all('div', class_='result')
        if results:
            output = []
            for i, res in enumerate(results[:3]):
                title_elem = res.find('a', class_='result__a')
                snippet_elem = res.find('a', class_='result__snippet')
                if title_elem:
                    title = title_elem.get_text(strip=True)
                    link = title_elem.get('href')
                    snippet = snippet_elem.get_text(strip=True) if snippet_elem else "Нет описания"
                    output.append(f"{i+1}. {title}\nСсылка: {link}\nОписание: {snippet[:200]}...")
            
            # Проверяем, не мусор ли это
            if output and not any("Microsoft" in o or "Mac" in o or "Poki" in o for o in output):
                return "\n\n".join(output)
    except Exception:
        pass

    # 4. Если поисковики заблокировали — даем кликабельную ссылку на Яндекс
    yandex_link = f"https://yandex.ru/search/?text={requests.utils.quote(clean_query)}"
    return f"❌ Поисковики заблокировали робота. Но вы можете открыть результаты вручную:\n\n🔗 <a href='{yandex_link}' target='_blank' style='color:#0284c7; font-weight:bold; font-size:16px;'>Нажмите сюда, чтобы открыть поиск в Яндексе</a>"

@app.get("/downloads/{filename}")
def download_agent_file(filename: str):
    file_path = os.path.join("downloads", filename)
    if os.path.exists(file_path):
        return FileResponse(file_path, filename=filename)
    return HTMLResponse("Файл не найден", status_code=404)

@app.get("/", response_class=HTMLResponse)
def read_root(request: Request):
    return templates.TemplateResponse(request=request, name="index.html")

# --- Остальные модули (сокращены для читаемости, но полностью функциональны) ---
@app.post("/calculate_duty", response_class=HTMLResponse)
def calculate_duty(request: Request, estate_value: float = Form(...), kinship: str = Form(...)):
    duty = min(estate_value * 0.003, 100000.0) if kinship == "close" else min(estate_value * 0.006, 1000000.0)
    return templates.TemplateResponse(request=request, name="index.html", context={"duty_result": f"{duty:,.2f}".replace(",", " ")})

@app.post("/calculate_mandatory_share", response_class=HTMLResponse)
def calculate_mandatory_share(request: Request, total_heirs: int = Form(...)):
    share = Fraction(1, total_heirs * 2)
    return templates.TemplateResponse(request=request, name="index.html", context={"share_result": f"{share.numerator}/{share.denominator}"})

@app.post("/upload_passport", response_class=HTMLResponse)
async def upload_passport(request: Request, file: UploadFile = File(...)):
    passport_data = None
    try:
        contents = await file.read()
        image = Image.open(io.BytesIO(contents)).convert('RGB')
        extracted_text = pytesseract.image_to_string(image, lang='rus+eng')
        match_mrz = re.search(r'(\d{4})\s*(\d{6})', extracted_text)
        if match_mrz:
            passport_data = {"series": f"{match_mrz.group(1)[:2]} {match_mrz.group(1)[2:]}", "number": match_mrz.group(2)}
    except Exception as e:
        extracted_text = f"Ошибка: {str(e)}"
    return templates.TemplateResponse(request=request, name="index.html", context={"ocr_result": extracted_text, "passport_data": passport_data})

@app.post("/save_client")
async def save_client(full_name: str = Form(...), series: str = Form(...), number: str = Form(...), phone: str = Form(""), email: str = Form("")):
    conn = sqlite3.connect('notary.db')
    conn.cursor().execute('INSERT INTO clients (full_name, series, number, phone, email) VALUES (?, ?, ?, ?, ?)', (full_name, series, number, phone, email))
    conn.commit()
    conn.close()
    return RedirectResponse(url="/journal", status_code=303)

@app.get("/journal", response_class=HTMLResponse)
def view_journal(request: Request):
    conn = sqlite3.connect('notary.db')
    clients = conn.cursor().execute('SELECT id, full_name, series, number, phone, email, created_at FROM clients ORDER BY id DESC').fetchall()
    conn.close()
    return templates.TemplateResponse(request=request, name="journal.html", context={"clients": clients})

@app.get("/download_docx/{client_id}")
def download_docx(client_id: int):
    conn = sqlite3.connect('notary.db')
    client = conn.cursor().execute('SELECT full_name, series, number FROM clients WHERE id = ?', (client_id,)).fetchone()
    conn.close()
    doc = Document()
    doc.add_heading('СОГЛАСИЕ НА ОБРАБОТКУ ПЕРСОНАЛЬНЫХ ДАННЫХ', 0)
    doc.add_paragraph(f'Я, {client[0]}, паспорт серии {client[1]} № {client[2]}, даю свое согласие на обработку моих персональных данных нотариусу.')
    doc.add_paragraph('\nДата: _______________________          Подпись: _______________________')
    file_stream = io.BytesIO(); doc.save(file_stream); file_stream.seek(0)
    return StreamingResponse(file_stream, headers={'Content-Disposition': f'attachment; filename="soglasie_{client_id}.docx"'}, media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

@app.post("/create_case")
def create_case(client_id: int = Form(...), case_number: str = Form(...), description: str = Form(...)):
    conn = sqlite3.connect('notary.db')
    conn.cursor().execute('INSERT INTO cases (client_id, case_number, description) VALUES (?, ?, ?)', (client_id, case_number, description))
    conn.commit()
    conn.close()
    return RedirectResponse(url="/cases", status_code=303)

@app.get("/cases", response_class=HTMLResponse)
def view_cases(request: Request):
    conn = sqlite3.connect('notary.db')
    query = '''SELECT c.id, c.case_number, cl.full_name, c.description, c.status, c.created_at, c.client_id 
               FROM cases c JOIN clients cl ON c.client_id = cl.id ORDER BY c.id DESC'''
    cases = conn.cursor().execute(query).fetchall()
    clients = conn.cursor().execute('SELECT id, full_name FROM clients').fetchall()
    conn.close()
    return templates.TemplateResponse(request=request, name="cases.html", context={"cases": cases, "clients": clients})

@app.get("/case/{case_id}", response_class=HTMLResponse)
def case_detail(request: Request, case_id: int):
    conn = sqlite3.connect('notary.db')
    query = '''SELECT c.id, c.case_number, c.description, c.status, c.created_at, 
                      cl.full_name, cl.series, cl.number, cl.phone, cl.email
               FROM cases c JOIN clients cl ON c.client_id = cl.id WHERE c.id = ?'''
    case_data = conn.cursor().execute(query, (case_id,)).fetchone()
    conn.close()
    if not case_data: return HTMLResponse("Дело не найдено")
    case = {"id": case_data[0], "number": case_data[1], "description": case_data[2], "status": case_data[3], "date": case_data[4][:10], "client_name": case_data[5], "series": case_data[6], "number_pass": case_data[7], "phone": case_data[8], "email": case_data[9]}
    return templates.TemplateResponse(request=request, name="case_detail.html", context={"case": case})

@app.post("/api/case_chat/{case_id}")
async def api_case_chat(request: Request, case_id: int):
    data = await request.json()
    user_message = data.get("message", "")
    conn = sqlite3.connect('notary.db')
    query = '''SELECT c.case_number, c.description, cl.full_name, cl.series, cl.number FROM cases c JOIN clients cl ON c.client_id = cl.id WHERE c.id = ?'''
    case_info = conn.cursor().execute(query, (case_id,)).fetchone()
    conn.close()
    if not case_info: return {"response": "Дело не найдено"}
    prompt = f"""<|im_start|>system
Ты — персональный ИИ-помощник. ФИО: {case_info[2]}, Паспорт: {case_info[3]} {case_info[4]}. Дело: {case_info[0]}.
Отвечай на русском языке. КАТЕГОРИЧЕСКИ ЗАПРЕЩЕНО использовать английские шаблоны.
<|im_end|>
<|im_start|>user
{user_message}
<|im_end|>
<|im_start|>assistant
"""
    try:
        response = requests.post('http://localhost:11434/api/generate', json={"model": "qwen2.5", "prompt": prompt, "stream": False, "options": {"temperature": 0.1, "num_thread": 8, "num_predict": 500}}, timeout=120)
        ai_reply = clean_ai_text(response.json().get('response', 'Ошибка').replace("<|im_end|>", ""))
    except Exception: ai_reply = "Ошибка: ИИ (Ollama) не запущен."
    return {"response": ai_reply}

@app.post("/generate_case_doc/{case_id}")
def generate_case_doc(case_id: int, doc_type: str = Form(...)):
    conn = sqlite3.connect('notary.db')
    query = '''SELECT cl.full_name, cl.series, cl.number FROM cases c JOIN clients cl ON c.client_id = cl.id WHERE c.id = ?'''
    client = conn.cursor().execute(query, (case_id,)).fetchone()
    conn.close()
    if not client: return HTMLResponse("Клиент не найден")
    prompt = f"""<|im_start|>system
Ты — робот-генератор. Выведи ТОЛЬКО готовый текст документа. Начни сразу с "Уважаемый" или "Просим". 
КАТЕГОРИЧЕСКИ ЗАПРЕЩЕНО писать вступления, вопросы или комментарии. Только русский язык.
Клиент (ФИО: {client[0]}, Паспорт: {client[1]} {client[2]}) — это НАСЛЕДНИК. Пиши от имени нотариуса.
Если не хватает данных, ставь ровную линию: "___________________".
<|im_end|>
<|im_start|>user
Составь документ: {doc_type}.
<|im_end|>
<|im_start|>assistant
"""
    try:
        response = requests.post('http://localhost:11434/api/generate', json={"model": "qwen2.5", "prompt": prompt, "stream": False, "options": {"temperature": 0.1, "num_thread": 8, "num_predict": 600}}, timeout=120)
        ai_text = clean_ai_text(response.json().get('response', 'Ошибка генерации').replace("<|im_end|>", ""))
    except Exception: ai_text = "Ошибка: Локальный ИИ не запущен."

    doc = Document()
    style = doc.styles['Normal']; style.font.name = 'Times New Roman'; style.font.size = Pt(14)
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f'Документ: {doc_type}'); run.bold = True
    for paragraph in ai_text.split('\n'):
        if paragraph.strip():
            p = doc.add_paragraph(paragraph.strip())
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(1.25)
    file_stream = io.BytesIO(); doc.save(file_stream); file_stream.seek(0)
    return StreamingResponse(file_stream, headers={'Content-Disposition': f'attachment; filename="doc_{case_id}.docx"'}, media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

# --- МОДУЛЬ: МУЛЬТИЗАДАЧНЫЙ AI-АГЕНТ ---
@app.get("/agent", response_class=HTMLResponse)
def agent_page(request: Request):
    return templates.TemplateResponse(request=request, name="agent.html")

@app.post("/api/agent_execute", response_class=HTMLResponse)
def agent_execute(request: Request, task: str = Form(...)):
    action_log = []
    
    prompt = f"""<|im_start|>system
Ты — AI-Агент. Выбери нужный инструмент для задачи.
Инструменты:
1. "create_word_doc": title (название), content (текст).
2. "read_website": url (ссылка).
3. "calculate_duty": value (число), kinship ("close" или "other").
4. "list_clients": без аргументов.
5. "search_web": query (поисковый запрос).

ЖЕСТКИЕ ПРАВИЛА:
1. Если пользователь не дал тебе явную ссылку (начинающуюся на http), ты ОБЯЗАН использовать "search_web".
2. НИКОГДА не используй "read_website", если ссылки нет в тексте запроса.
3. Выведи СТРОГО JSON-массив действий. Пример: [{{"tool": "search_web", "query": "нотариальная палата Волгограда"}}]
Никакого другого текста, только JSON.
<|im_end|>
<|im_start|>user
{task}
<|im_end|>
<|im_start|>assistant
"""
    try:
        response = requests.post('http://localhost:11434/api/generate', json={
            "model": "qwen2.5", "prompt": prompt, "stream": False, 
            "options": {"temperature": 0.1, "num_thread": 8, "num_predict": 500}
        }, timeout=60) 
        
        ai_response = response.json().get('response', '').replace("<|im_end|>", "").strip()
        ai_response = ai_response.replace('```json', '').replace('```', '').strip()
        
        try:
            actions = json.loads(ai_response)
            if not isinstance(actions, list): actions = [actions]
            
            for act in actions:
                tool = act.get("tool")
                
                if tool == "search_web":
                    query = act.get("query", "")
                    action_log.append(f"🔍 Ищу в интернете: {query}...")
                    search_results = search_web(query)
                    action_log.append(f"📝 Результаты:\n{search_results}")

                elif tool == "list_clients":
                    conn = sqlite3.connect('notary.db')
                    clients = conn.cursor().execute('SELECT full_name, series, number FROM clients').fetchall()
                    conn.close()
                    action_log.append(f"👥 Список клиентов в базе: {len(clients)} чел.")
                    for c in clients:
                        action_log.append(f" - {c[0]} (паспорт {c[1]} {c[2]})")
                
                elif tool == "calculate_duty":
                    val = float(act.get("value", 0))
                    kin = act.get("kinship", "close")
                    duty = min(val * 0.003, 100000.0) if kin == "close" else min(val * 0.006, 1000000.0)
                    action_log.append(f"🧮 Госпошлина за имущество {val:,.2f} руб.: {duty:,.2f} руб.")

                elif tool == "read_website":
                    url = act.get("url", "")
                    
                    if url not in task and "http" not in task:
                        action_log.append(f"⚠️ Агент попытался выдумать ссылку ({url}). Останавливаю и запускаю поиск...")
                        search_results = search_web(task)
                        action_log.append(f"📝 Результаты поиска:\n{search_results}")
                        action_log.append("💡 Выберите нужную ссылку и скопируйте её в новый запрос.")
                        continue 

                    action_log.append(f"🌐 Пытаюсь прочитать сайт: {url}...")
                    try:
                        web_resp = requests.get(url, timeout=10, verify=False, headers={'User-Agent': 'Mozilla/5.0'})
                        web_resp.encoding = 'utf-8'
                        soup = BeautifulSoup(web_resp.text, 'html.parser')
                        for script in soup(["script", "style"]): script.extract()
                        text = soup.get_text(separator=' ', strip=True)
                        
                        if "404" in text or "Not Found" in text or len(text) < 100:
                            action_log.append(f"⚠️ Сайт {url} вернул ошибку 404. Ищу правильный адрес...")
                            search_results = search_web(task)
                            action_log.append(f"🔍 Результаты поиска:\n{search_results}")
                        else:
                            action_log.append(f"📄 Удалось прочитать (фрагмент):\n{text[:1000]}...")
                    except Exception as e:
                        action_log.append(f"❌ Ошибка чтения: {str(e)}")

                elif tool == "create_word_doc":
                    title = act.get("title", "Документ")
                    content = act.get("content", "Нет текста")
                    doc = Document()
                    style = doc.styles['Normal']; style.font.name = 'Times New Roman'; style.font.size = Pt(14)
                    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = t.add_run(title); r.bold = True
                    for p in content.split('\n'):
                        if p.strip():
                            par = doc.add_paragraph(p.strip())
                            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                            par.paragraph_format.first_line_indent = Cm(1.25)
                    filename = f"agent_doc_{int(time.time())}.docx"
                    filepath = os.path.join("downloads", filename)
                    doc.save(filepath)
                    action_log.append(f"📄 Создан документ: <a href='/downloads/{filename}' target='_blank' style='color:#0284c7; font-weight:bold;'>⬇️ Скачать {title}</a>")

        except json.JSONDecodeError:
            action_log.append("⚠️ ИИ не смог составить план. Попробуйте переформулировать задачу проще.")
            action_log.append(f"Ответ ИИ: {ai_response}")

    except requests.exceptions.Timeout:
        action_log.append("⏳ Ошибка: Локальный ИИ (Ollama) слишком долго думает. Попробуйте еще раз.")
    except Exception as e:
        action_log.append(f"❌ Ошибка связи с ИИ: {str(e)}")

    result_html = "<br>".join(action_log)
    return templates.TemplateResponse(request=request, name="agent.html", context={"result": result_html, "task": task})

# ОСТАЛЬНЫЕ МОДУЛИ (ЕГРН, Ревизор, Чат)
@app.post("/analyze_egrn/{case_id}")
async def analyze_egrn(case_id: int, file: UploadFile = File(...)):
    try:
        contents = await file.read()
        image = Image.open(io.BytesIO(contents)).convert('RGB')
        raw_text = pytesseract.image_to_string(image, lang='rus+eng')
        prompt = f"""<|im_start|>system
Ты — робот-парсер недвижимости. Извлеки: Адрес, Кадастровый номер, Площадь, Стоимость, ФИО собственника. Верни СТРОГО JSON.
<|im_end|>
<|im_start|>user
Текст документа:
{raw_text}
<|im_end|>
<|im_start|>assistant
"""
        response = requests.post('http://localhost:11434/api/generate', json={"model": "qwen2.5", "prompt": prompt, "stream": False, "options": {"temperature": 0.1, "num_thread": 8, "num_predict": 200}}, timeout=120)
        ai_text = response.json().get('response', '{}').strip().replace("```json", "").replace("```", "").strip()
        try: data = json.loads(ai_text)
        except: data = {"error": "Ошибка распознавания"}
        doc = Document()
        style = doc.styles['Normal']; style.font.name = 'Times New Roman'; style.font.size = Pt(14)
        doc.add_heading('РЕЗЮМЕ ПО ОБЪЕКТУ НЕДВИЖИМОСТИ', 0)
        if "error" in data: doc.add_paragraph(data["error"])
        else:
            value_text = str(data.get('value', 'Не указано')).replace('руб.', '').replace('руб', '').strip()
            doc.add_paragraph(f"📍 Адрес: {data.get('address', 'Не указано')}")
            doc.add_paragraph(f"🔢 Кадастровый номер: {data.get('cadastral', 'Не указано')}")
            doc.add_paragraph(f"📐 Площадь: {data.get('area', 'Не указано')}")
            doc.add_paragraph(f"💰 Стоимость: {value_text} руб.")
            doc.add_paragraph(f"👤 Собственник: {data.get('owner', 'Не указано')}")
        file_stream = io.BytesIO(); doc.save(file_stream); file_stream.seek(0)
        return StreamingResponse(file_stream, headers={'Content-Disposition': f'attachment; filename="egrn_summary_{case_id}.docx"'}, media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    except Exception as e:
        return HTMLResponse(f"Ошибка: {str(e)}", status_code=500)

@app.get("/chat", response_class=HTMLResponse)
def chat_page(request: Request):
    return templates.TemplateResponse(request=request, name="chat.html")

@app.post("/api/chat")
async def api_chat(request: Request):
    data = await request.json()
    prompt = f"""<|im_start|>system
Ты — опытный российский нотариус. Отвечай на вопросы. Запрещены китайские иероглифы и английский язык.
<|im_end|>
<|im_start|>user
{data.get("message", "")}
<|im_end|>
<|im_start|>assistant
"""
    try:
        response = requests.post('http://localhost:11434/api/generate', json={"model": "qwen2.5", "prompt": prompt, "stream": False, "options": {"temperature": 0.2, "num_thread": 8, "num_predict": 500}}, timeout=120)
        ai_reply = clean_ai_text(response.json().get('response', 'Ошибка').replace("<|im_end|>", ""))
    except Exception: ai_reply = "Ошибка: ИИ не запущен."
    return {"response": ai_reply}

@app.get("/checker", response_class=HTMLResponse)
def checker_page(request: Request):
    return templates.TemplateResponse(request=request, name="checker.html")

@app.post("/api/check_doc", response_class=HTMLResponse)
async def check_doc(request: Request, file: UploadFile = File(...)):
    global last_checker_result
    try:
        contents = await file.read()
        image = Image.open(io.BytesIO(contents)).convert('RGB')
        raw_text = pytesseract.image_to_string(image, lang='rus+eng')
        prompt = f"""<|im_start|>system
Ты — строгий юридический ревизор. Найди до 5 критичных ошибок. Запрещены китайские иероглифы и английский язык.
<|im_end|>
<|im_start|>user
Текст документа:
{raw_text}
<|im_end|>
<|im_start|>assistant
"""
        response = requests.post('http://localhost:11434/api/generate', json={"model": "qwen2.5", "prompt": prompt, "stream": False, "options": {"temperature": 0.1, "num_thread": 8, "num_predict": 800}}, timeout=120)
        ai_text = clean_ai_text(response.json().get('response', 'Ошибка анализа').replace("<|im_end|>", ""))
        last_checker_result["text"] = ai_text
    except Exception as e:
        ai_text = f"Ошибка: {str(e)}"
        last_checker_result["text"] = ai_text
    return templates.TemplateResponse(request=request, name="checker.html", context={"result": ai_text})

@app.get("/download_checker_report")
def download_checker_report():
    if not last_checker_result["text"]: return HTMLResponse("Отчет пуст.")
    doc = Document()
    style = doc.styles['Normal']; style.font.name = 'Times New Roman'; style.font.size = Pt(14)
    doc.add_heading('ЗАКЛЮЧЕНИЕ ИИ-РЕВИЗОРА', 0)
    for paragraph in last_checker_result["text"].split('\n'):
        if paragraph.strip():
            p = doc.add_paragraph(paragraph.strip())
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(1.25)
    file_stream = io.BytesIO(); doc.save(file_stream); file_stream.seek(0)
    return StreamingResponse(file_stream, headers={'Content-Disposition': 'attachment; filename="checker_report.docx"'}, media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')